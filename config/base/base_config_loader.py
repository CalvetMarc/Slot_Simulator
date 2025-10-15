import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

TABLE_NAMES = ["Grid", "Strips", "Paylines", "Paytable"]

def load_all_tables(base_path, file_name="slot_config.xlsx"):
    """
    Loads all standard tables (Grid, Strips, Paylines, Paytable) from the Excel file
    and returns them as pandas DataFrames.
    """
    excel_file = Path(base_path) / file_name
    wb = load_workbook(excel_file, data_only=True)
    result = {}

    for table_name in TABLE_NAMES:
        df = None
        for ws in wb.worksheets:
            if table_name in ws.tables:
                ref = ws.tables[table_name].ref
                start_cell, end_cell = ref.split(":")
                start_col = ''.join([c for c in start_cell if c.isalpha()])
                start_row = int(''.join([c for c in start_cell if c.isdigit()]))
                end_col = ''.join([c for c in end_cell if c.isalpha()])
                end_row = int(''.join([c for c in end_cell if c.isdigit()]))

                df = pd.read_excel(
                    excel_file,
                    sheet_name=ws.title,
                    header=None,
                    usecols=f"{start_col}:{end_col}",
                    skiprows=start_row - 1,
                    nrows=end_row - start_row + 1
                )

                df.columns = df.iloc[0]
                df = df[1:].reset_index(drop=True)
                result[table_name.lower()] = df
                break

        if df is None:
            print(f"⚠️ Table '{table_name}' not found.")

    return result
